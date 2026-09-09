"""Read-only reassessment of the preserved 2020 D timed trial.

Run with --workspace pointing to the original corpus. Outputs go to a new
directory; the old score, timeline, source files and paper remain unchanged.
"""
import argparse
import hashlib
import json
import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "skill/cumcm-reliable-paper/scripts"))
from curve_coverage import directed_coverage
from paper_structure import check_structure


def load(path):
    return json.loads(path.read_text(encoding="utf-8"))


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def audit(workspace, output):
    legacy = workspace / "blind_tests/2020_D_profilometer_strict"
    paper = legacy / "output/pdf/2020D_profilometer_strict_blind_test_paper.pdf"
    result_path = legacy / "output/results.json"
    result = load(result_path)
    recorded = load(legacy / "output/final_audit.json")
    if digest(paper) != recorded["artifacts"]["pdf_sha256"] or digest(result_path) != recorded["artifacts"]["results_sha256"]:
        raise ValueError("legacy audited paper or results have changed")
    source_manifest = load(legacy / "source_manifest.json")
    for source in source_manifest["sources"]:
        if digest(workspace / source["path"]) != source["sha256"]:
            raise ValueError("official source hash mismatch")
    output.mkdir(parents=True, exist_ok=False)
    timeline = load(legacy / "timeline.json")
    pages = [page.extract_text() for page in PdfReader(paper).pages]
    text = "\n".join(pages)
    w2 = result["workpiece2"]
    contour_path = legacy / w2["merged_contour"]["ordered_path"]
    contour = np.loadtxt(contour_path, delimiter=",", skiprows=1)[:, 1:3]
    tilts = {row["sheet"]: row["estimated_deg"] for row in w2["tilts"]}
    shifts = {row["sheet"]: np.asarray(row["shift"]) for row in w2["registrations"]}
    shifts[w2["reference_sheet"]] = np.zeros(2)
    workbook = workspace / "2020/D/附件2_工件2的整体测量数据.xlsx"
    wb = load_workbook(workbook, read_only=True, data_only=True)
    coverage = []
    try:
        for sheet in wb.worksheets:
            raw = np.asarray(list(sheet.iter_rows(min_row=2, max_col=2, values_only=True)), float)
            angle = np.deg2rad(-tilts[sheet.title])
            rotation = np.array([[np.cos(angle), -np.sin(angle)], [np.sin(angle), np.cos(angle)]])
            aligned = raw @ rotation.T + shifts[sheet.title]
            # Exactly -20 is the original model's censoring hypothesis, not a
            # verified instrument specification. Do not connect across masked gaps.
            retained = raw[:, 1] != -20
            edges = np.diff(np.r_[False, retained, False].astype(int))
            segments = [(i, j) for i, j in zip(np.flatnonzero(edges == 1), np.flatnonzero(edges == -1)) if j - i >= 2]
            reports = [directed_coverage(aligned[i:j], contour, tolerance=.03, step=.005) for i, j in segments]
            if not reports:
                raise ValueError("no non-boundary source segment")
            length = sum(r["source_length"] for r in reports)
            coverage.append({
                "sheet": sheet.title, "raw_points": len(raw), "boundary_points": int((~retained).sum()),
                "non_boundary_source_length": length,
                "covered_fraction_lower": sum(r["source_length"] * r["covered_fraction_lower"] for r in reports) / length,
                "covered_fraction_upper": sum(r["source_length"] * r["covered_fraction_upper"] for r in reports) / length,
                "maximum_distance_lower": max(r["maximum_distance_lower"] for r in reports),
                "maximum_distance_upper": max(r["maximum_distance_upper"] for r in reports),
                "segments": reports,
            })
            print({k: coverage[-1][k] for k in ("sheet", "covered_fraction_upper", "maximum_distance_lower")}, flush=True)
    finally:
        wb.close()
    structure = check_structure(paper, question_count=4)
    source = (legacy / "analysis.py").read_text(encoding="utf-8")
    # These observed anchors identify this particular preserved failure; they
    # are not a generic semantic paper checker.
    if not ("选择物理" in pages[0] and "不参与估计" in text and "tilt, candidates = estimate_tilt(raw)" in source):
        raise ValueError("review anchors changed; reassess prose/code consistency")
    observed = {
        "case": "2020 D preserved 180-minute trial",
        "decision": "NOT_ACCEPTED_AS_COMPLETE_TIMED_PAPER",
        "review_type": "agent reassessment, not external expert or award judgment",
        "historical_record": {"claimed_score": 93, "claimed_p0_failures": 0,
                              "recorded_elapsed_seconds": timeline["events"][-1]["elapsed_seconds"],
                              "time_scope": "local timeline reports duration; this audit cannot independently attest original read isolation or contemporaneous artifact existence"},
        "paper_structure": structure,
        "coverage": coverage,
        "coverage_scope": "uses saved transforms and supplied ordered contour; excludes original z=-20 points without joining gaps; .03 is the legacy diagnostic threshold, not an official measurement tolerance; raw noise contributes to arc length; coverage is not registration truth",
        "findings": [
            {"id": "PAPER_FORMAT", "severity": "blocking", "evidence": "Page 1 includes abstract, restatement, assumptions and methods; formula subscripts and sqrt are printed as prose. No dedicated per-question analysis or notation table.", "required_resolution": "A readable Chinese mathematical paper with the requested structure, reviewed on rendered pages."},
            {"id": "METHOD_DRIFT", "severity": "blocking", "evidence": "Page 1 section 3.2 says sheet names choose angle branches; page 3 says they are only checked afterward. analysis.py calls estimate_tilt(raw) without the nominal name.", "required_resolution": "Reconcile actual estimation, independent validation and all paper descriptions."},
            {"id": "Q3_OUTPUTS", "severity": "blocking", "evidence": "Official Q3 requests all Q1 parameter types and complete contour. The delivered Q3 summary contains tilts, one primary circle and a reference-backbone consensus. It does not enumerate all required feature measurements.", "required_resolution": "Deliver the full feature inventory with observed/estimated/unresolved statuses; justify any irrecoverable item from coverage and feature correspondence evidence."},
            {"id": "Q3_UNION", "severity": "blocking", "evidence": "analysis.py constructs consensus only at resampled reference points. Other scans contribute only within a .03 tube. This cannot establish the union of all measured regions; full directed coverage is reported here.", "required_resolution": "Integrate supported extensions or explicitly limit the output to the reference domain and quantify omitted observations."},
            {"id": "Q4_CORRECTION", "severity": "blocking", "evidence": "main() produces the contour before local circle/angle analysis; corrected_workpiece2_parameters stores group summaries only. No geometric correspondence map or updated contour is generated afterward.", "required_resolution": "Map local features to global features, apply constrained corrections and output the resulting contour with before/after residuals."},
            {"id": "WEAK_ACCEPTANCE", "severity": "blocking", "evidence": "final_audit.py explicitly requires four pages and matches headings/numeric tokens. independent_check.py reads the original results.json for its 13 checks, rather than independently reconstructing Q3/Q4 from measurements.", "required_resolution": "Retain these as consistency checks; accept each official requested output on independent evidence and actual paper review."}
        ],
        "retained_evidence": "Source hashes and recorded duration agree with the archived audit; some geometric computations and consistency checks remain useful. No claim here that every old numerical estimate is wrong.",
        "source_sha256": {str(p.relative_to(workspace)): digest(p) for p in [paper, result_path, contour_path, legacy / "analysis.py", legacy / "geometry_core.py", legacy / "independent_check.py", legacy / "final_audit.py", legacy / "timeline.json", legacy / "INDEPENDENT_PAPER_REVIEW.md"]},
        "tool_sha256": {str(p.relative_to(ROOT)): digest(p) for p in [Path(__file__), ROOT / "skill/cumcm-reliable-paper/scripts/curve_coverage.py"]},
        "submission_ready": False, "national_award_level": "NOT_ESTABLISHED",
    }
    (output / "report.json").write_text(json.dumps(observed, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("Audit complete; historical acceptance not adopted.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workspace", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    audit(args.workspace.resolve(), args.output.resolve())
