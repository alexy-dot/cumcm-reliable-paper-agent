"""Read frozen official data and evaluate order-to-supply baselines chronologically."""
import argparse
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


def write(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def rates(o, s, method):
    if method == "volume_ratio":
        return np.divide(s.sum(1), o.sum(1), out=np.zeros(len(o)), where=o.sum(1)>0)
    if method == "least_squares":
        return np.divide((s*o).sum(1), (o*o).sum(1), out=np.zeros(len(o)), where=(o*o).sum(1)>0)
    if method == "median_ratio":
        return np.array([np.median(b[a>0]/a[a>0]) if (a>0).any() else 0 for a,b in zip(o,s)])
    raise ValueError(method)


def read_data(run):
    p = next((run/"sources").glob("*附件1*"))
    w = load_workbook(p, read_only=True, data_only=True)
    rows = [list(s.iter_rows(min_row=2, max_row=403, values_only=True)) for s in w]
    w.close()
    ids=[r[0] for r in rows[0]]; kinds=[r[1] for r in rows[0]]
    assert len(ids)==402 and len(set(ids))==402
    assert [(r[0],r[1]) for r in rows[0]]==[(r[0],r[1]) for r in rows[1]]
    o,s=[np.asarray([r[2:] for r in rr],float) for rr in rows]
    assert o.shape==s.shape==(402,240) and np.isfinite(o).all() and np.isfinite(s).all()
    assert (o>=0).all() and (s>=0).all() and not ((o==0)&(s>0)).any()
    p2=next((run/"sources").glob("*附件2*"));w=load_workbook(p2,read_only=True,data_only=True)
    rows=list(w.active.iter_rows(min_row=2,values_only=True));w.close()
    losses=np.asarray([r[1:] for r in rows],float)/100
    assert losses.shape==(8,240) and np.isfinite(losses).all() and (losses>=0).all() and (losses<1).all()
    return ids,kinds,o,s,losses


def prepare(run):
    ids,kinds,o,s,losses=read_data(run)
    contract=json.loads((run/"problem_contract.json").read_text())
    contract["problem_interpretation"]="在供货不完全遵循订货、转运存在损耗的条件下，评价402家供应商，制定24周订购与转运计划，并分析材料结构与增产。"
    contract["official_outputs"]=["50家重要供应商名单", "问题2、3、4未来24周订购和转运方案", "保持附件A、附件B原文件名和已有信息，仅填写规定区域"]
    tasks=[("量化402家供应商供货特征并确定最重要的50家",["量化指标与模型", "论文中50家供应商名单"]),
           ("确定满足生产需求至少需多少家供应商，并给出最经济订购及最少损耗转运方案",["最少供应商数量", "24周订购与转运方案", "实施效果分析"]),
           ("尽量多采购A、少采购C，并降低转运损耗",["新订购与转运方案", "实施效果分析"]),
           ("根据现有供应商和转运商确定每周可提高的产能",["产能提升量", "相应24周订购和转运方案"])]
    contract["questions"]=[{"id":f"Q{i}","source_locator":f"官方题面第1页问题{i}及第2页附件说明", "objective":task,
        "inputs":["附件1完整402×240订货与供货数据", "附件2完整8×240损耗记录"],"outputs":outputs,
        "decision_variables":[] if i==1 else ["供应商订货量", "供应商—转运商分配量"]+(["每周产能"] if i==4 else []),
        "hard_constraints":["每家转运商每周运输量不超过6000立方米", "实际提供的原材料全部收购", "A/B/C每单位产品耗材0.6/0.66/0.72立方米"] if i>1 else [],
        "units":["原材料立方米", "产品立方米", "周", "损耗率为百分数转换后的比例"],
        "ambiguities":["原题未给初始库存和绝对运储单价", "尽可能两周库存与尽量单家转运为柔性表述；建模选择须另作说明"],
        "termination_condition":"完成该问的全部指定输出并验证"} for i,(task,outputs) in enumerate(tasks,1)]
    write(run/"problem_contract.json",contract)
    a=np.array([{"A":.6,"B":.66,"C":.72}[k] for k in kinds])
    price=np.array([{"A":1.2,"B":1.1,"C":1.}[k] for k in kinds])
    methods=("volume_ratio","least_squares","median_ratio")
    comparison=[]
    for method in methods:
        r=rates(o[:,:144],s[:,:144],method);pred=r[:,None]*o[:,144:192]
        comparison.append({"method":method,"validation_active_mae":float(np.abs(pred-s[:,144:192])[o[:,144:192]>0].mean()),
                           "validation_total_weekly_mae":float(np.abs(pred.sum(0)-s[:,144:192].sum(0)).mean())})
    # The planning estimand is expected volume, not conditional median volume.
    # Preserve the MAE pilot as a diagnostic, never select a median as a mean.
    chosen="volume_ratio"
    r_train=rates(o[:,:192],s[:,:192],chosen);pred=r_train[:,None]*o[:,192:]
    active=o[:,192:]>0
    test={"weeks":"193-240", "mean_accounting_method":chosen,
          "active_mae":float(np.abs(pred-s[:,192:])[active].mean()),
          "weekly_total_mae":float(np.abs(pred.sum(0)-s[:,192:].sum(0)).mean()),
          "actual_weekly_total_mean":float(s[:,192:].sum(0).mean()),
          "scope":"chronological diagnostic at observed historical orders; median pilot exposed this block before the estimand correction, so this is not an untouched final test or a counterfactual evaluation of procurement"}
    fitted=rates(o,s,chosen)
    # Standard weekly contract capacity uses a historical active-week mean.
    # No-order weeks are not evidence that the supplier was unavailable.
    caps=np.array([row[row>0].mean() if (row>0).any() else 0 for row in o])
    caps90=np.array([np.quantile(row[row>0],.9) if (row>0).any() else 0 for row in o])
    supply_caps=caps*fitted
    carrier_mean=np.array([v[v>0].mean() for v in losses])
    carrier_p95=np.array([np.quantile(v[v>0],.95) for v in losses])
    equivalent=s/a[:,None];total=equivalent.sum(0);demand=28200.
    impact=np.minimum(demand,total)[None,:]-np.minimum(demand,total[None,:]-equivalent)
    score=impact.mean(1)
    order=np.lexsort((np.arange(402),-score))
    ranking=[]
    for rank,i in enumerate(order,1):
        mask=o[i]>0
        ranking.append({"rank":rank,"supplier":ids[i],"kind":kinds[i],"marginal_weekly_product_equivalent":float(score[i]),
                        "mean_supply":float(s[i].mean()),"ordered_weeks":int(mask.sum()),
                        "zero_supply_given_order":float((s[i,mask]==0).mean()) if mask.any() else None,
                        "volume_fulfilment":float(s[i].sum()/o[i].sum()) if mask.any() else None,
                        "order_cap":float(caps[i]),"fitted_supply_rate":float(fitted[i]),"modeled_supply_cap":float(supply_caps[i])})
    basic={"ids":ids,"kinds":kinds,"conversion":a.tolist(),"price":price.tolist(),"demand":demand,
           "order_cap":caps.tolist(),"order_cap_q90":caps90.tolist(),"supply_rate":fitted.tolist(),"supply_cap":supply_caps.tolist(),
           "carrier_mean_loss":carrier_mean.tolist(),"carrier_p95_loss":carrier_p95.tolist(),"top50_indices":order[:50].tolist(),
           "initial_inventory_product_equivalent":2*demand,"weeks":24,
           "assumptions":["Initial inventory equals two weeks of planned output; arrivals are usable in the same week; no lead time is specified.",
                          "Post-production stock target is two weeks. Stationary expected plans keep this stock; realized risks are evaluated separately.",
                          "Expected supply is a nonnegative linear response calibrated by volume ratio, within mean-positive-order contract caps; repeating historical active-week activity is a modeling assumption, not proven available capacity.",
                          "Prices are indexed to C purchase price=1. Actual transport/storage unit prices are absent; report physical volume and stock rather than invented monetary totals.",
                          "All modeled supplied raw material is purchased and transported; a carrier can carry at most 6000 raw cubic meters weekly.",
                          "Split shipments are permitted by the soft wording; number of splits will be reduced and reported.",
                          "Min supplier count and capacity improvement are conditional on fitted supply envelopes, not guaranteed future physical capacities."]}
    stats={"suppliers":402,"weeks":240,"positive_orders":int((o>0).sum()),"zero_supply_given_order_count":int(((o>0)&(s==0)).sum()),
           "nonzero_supply_without_order":int(((o==0)&(s>0)).sum()),"loss_zero_missing_counts":(losses==0).sum(1).tolist(),
           "carrier_mean_loss":carrier_mean.tolist(),"carrier_p95_loss":carrier_p95.tolist(),
           "train_weeks":"1-144","validation_weeks":"145-192","model_comparison":comparison,"holdout":test,
           "total_modeled_raw_supply_cap":float(supply_caps.sum()),"total_modeled_product_cap_before_transport":float((supply_caps/a).sum()),
           "top50_overlap_with_mean_supply_equivalent":int(len(set(order[:50])&set(np.argsort(-equivalent.mean(1))[:50]))),
           "source_sha256":{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in (run/"sources").iterdir() if p.is_file()}}
    write(run/"artifacts/facts.json",basic);write(run/"artifacts/data-analysis.json",stats);write(run/"artifacts/ranking.json",ranking)
    np.savez_compressed(run/"artifacts/observations.npz",orders=o,supply=s,loss=losses)
    print(json.dumps(stats,ensure_ascii=False,indent=2))


if __name__=="__main__":
    ap=argparse.ArgumentParser();ap.add_argument("run",type=Path);prepare(ap.parse_args().run)
