"""Independent source accounting, analytic capacity bound and declared stress cases."""
import argparse
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from prepare import write


def greedy_capacity(cap,conversion,loss):
    """Rearrangement bound: best conversion first, then least-loss carriers.

    With unrestricted splitting and no other links forbidden, swapping two
    allocations improves yield by (1/a-1/b)*(loss_bad-loss_good)>=0.
    """
    remaining=[6000.]*8;total=0.;allocation=np.zeros((len(cap),8))
    for i in sorted(range(len(cap)),key=lambda i:conversion[i]):
        amount=cap[i]
        for j in sorted(range(8),key=lambda j:loss[j]):
            used=min(amount,remaining[j]);amount-=used;remaining[j]-=used
            allocation[i,j]+=used;total+=used*(1-loss[j])/conversion[i]
            if amount<1e-10:break
    return total,allocation


def verify(run):
    f=json.loads((run/"artifacts/facts.json").read_text());plans=json.loads((run/"artifacts/plans.json").read_text())
    wb=load_workbook(next((run/"sources").glob("*附件1*")),read_only=True,data_only=True)
    sheets=[list(s.iter_rows(min_row=2,max_row=403,values_only=True)) for s in wb];wb.close()
    orders=np.array([r[2:] for r in sheets[0]],float);supply=np.array([r[2:] for r in sheets[1]],float)
    kinds=[r[1] for r in sheets[0]];conv=[{"A":.6,"B":.66,"C":.72}[k] for k in kinds];price=[{"A":1.2,"B":1.1,"C":1.}[k] for k in kinds]
    cap=np.array([sum(s)/sum(v>0 for v in o) if any(o) else 0 for o,s in zip(orders,supply)])
    rate=np.array([sum(s)/sum(o) if sum(o)>0 else 0 for o,s in zip(orders,supply)])
    wb=load_workbook(next((run/"sources").glob("*附件2*")),read_only=True,data_only=True)
    loss=[sum(v for v in row[1:] if v>0)/sum(v>0 for v in row[1:])/100 for row in wb.active.iter_rows(min_row=2,values_only=True)];wb.close()
    checks=[];stress={}
    def add(name,error,tolerance):checks.append({"name":name,"error":float(error),"tolerance":tolerance,"passed":bool(error<=tolerance)})
    add("source-derived supply envelopes",np.max(np.abs(cap-np.array(f["supply_cap"]))),1e-8)
    add("source-derived conditional loss means",np.max(np.abs(np.array(loss)-f["carrier_mean_loss"])),1e-12)
    lower_cap=sorted((v*(1-min(loss))/a for v,a in zip(cap,conv)),reverse=True)
    n=plans["Q2"]["active_suppliers"]
    add("n-1 suppliers impossible even at best carrier loss",max(0.,sum(lower_cap[:n-1])-f["demand"]+1e-6),0.)
    for name,p in plans.items():
        y=np.array(p["shipments"]);q=np.array(p["orders"]);d=p["production_per_week"]
        add(name+" all expected supply transported",np.max(np.abs(q*rate-y.sum(2))),1e-7)
        add(name+" contract capacity",max(0.,np.max(y.sum(2)-cap)),1e-7)
        add(name+" carrier capacity",max(0.,np.max(y.sum(1))-6000),1e-7)
        add(name+" nonnegative shipments and orders",max(0.,-min(y.min(),q.min())),1e-8)
        inventory=2*f["demand"];inv=[];cost=0.;raw_loss=0.
        for t in range(24):
            delivered=0.
            for i in range(402):
                for j in range(8):
                    delivered+=y[t,i,j]*(1-loss[j])/conv[i]
                    cost+=y[t,i,j]*price[i];raw_loss+=y[t,i,j]*loss[j]
            inventory+=delivered-d;inv.append(inventory)
        add(name+" reserve floor",max(0.,2*d-min(inv)),1e-6)
        add(name+" inventory reconstruction",np.max(np.abs(np.array(inv)-p["ending_inventory_product_equivalent"])),1e-6)
        add(name+" purchase cost accounting",abs(cost-p["purchase_cost_index"]),1e-5)
        add(name+" physical transport loss accounting",abs(raw_loss-p["expected_transport_loss"]),1e-5)
        scenarios=[]
        for multiplier in (.8,.9,1.,1.1):
            received=(y*multiplier*(1-np.array(loss))[None,None,:]/np.array(conv)[None,:,None]).sum((1,2))
            inv2=2*f["demand"]+np.cumsum(received-d)
            scenarios.append({"supply_multiplier":multiplier,"terminal_inventory":float(inv2[-1]),
                              "reserve_shortfall_weeks":int((inv2<2*d-1e-6).sum()),
                              "first_production_shortage_week":next((i+1 for i,v in enumerate(inv2) if v<0),None),
                              "fixed_carrier_overload_weeks":int((y.sum(1)*multiplier>6000+1e-6).any(1).sum())})
        stress[name]=scenarios
    bound,_=greedy_capacity(cap,conv,loss)
    add("Q4 steady capacity independent rearrangement construction",abs(bound-plans["Q4"]["steady_capacity"]),1e-6)
    immediate=min(bound,(bound+2*f["demand"])/3)
    add("Q4 first-week inventory upper bound attained",abs(immediate-plans["Q4"]["production_per_week"]),1e-6)
    report={"passed":all(c["passed"] for c in checks),"checks":checks,
            "minimum_count_proof":{"count":n,"best_possible_output_with_n_minus_one":sum(lower_cap[:n-1]),"required":f["demand"],"scope":"conditional on the declared supplier envelopes"},
            "capacity_bound":{"greedy_steady":bound,"immediate_uniform":immediate,"proof":"rearrangement for divisible raw material plus first-week inventory inequality"},
            "stress":stress,"stress_scope":"deterministic response perturbations, not estimated event probabilities; negative inventory denotes unmet production, never physical negative stock",
            "plans_sha256":hashlib.sha256((run/"artifacts/plans.json").read_bytes()).hexdigest(),"source_recomputed":True}
    write(run/"artifacts/verification.json",report)
    if not report["passed"]:raise ValueError([c for c in checks if not c["passed"]])
    print({"checks":len(checks),"passed":True,"capacity_bound":report["capacity_bound"],"minimum_count_proof":report["minimum_count_proof"]},flush=True)


if __name__=="__main__":
    ap=argparse.ArgumentParser();ap.add_argument("run",type=Path);verify(ap.parse_args().run)
