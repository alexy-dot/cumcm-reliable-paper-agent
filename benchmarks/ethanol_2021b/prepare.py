"""Freeze original experimental inputs and a model-comparison protocol before fitting."""
import argparse
import csv
import math
import re
import sys
from pathlib import Path

import openpyxl

PROJECT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(PROJECT/"skill/cumcm-reliable-paper/scripts"))
from engine import initialize_run,read_json,write_json,sha256_file

HASHES={"CUMCM2021-B.pdf":"c2f542c34272a90e1cba2131a625171f951b51b15277d969f4340ce1e54865d7",
        "附件1.xlsx":"de8e67e2c3fafb703553cb20e46f1a6db7109c5039bd50fd450c1e093a7080a3",
        "附件2.xlsx":"9aa92011182076833ced7cc47b57c57875a85ed8fb2b8edd2b1f66dd3c10ac82"}
FEATURES=["temperature_c","co_mass_mg","hap_mass_mg","quartz_mass_mg","co_loading_pct","feed_ml_min","mode_II"]


def parse_formula(text):
    text=re.sub(r"\s+","",text)
    cobalt=re.search(r"([\d.]+)mg([\d.]+)wt%Co/SiO2",text)
    hap=re.search(r"(?:-|\+)([\d.]+)mgHAP",text)
    feed=re.search(r"乙醇浓度([\d.]+)ml/min",text)
    quartz=re.search(r"(?:-|\+)([\d.]+)mg石英砂",text)
    if cobalt is None or feed is None or (hap is None and "无HAP" not in text):
        raise ValueError("unrecognized catalyst description: "+text)
    return {"co_mass_mg":float(cobalt[1]),"co_loading_pct":float(cobalt[2]),
            "hap_mass_mg":float(hap[1]) if hap else 0.,"quartz_mass_mg":float(quartz[1]) if quartz else 0.,"feed_ml_min":float(feed[1])}


def extract(source):
    workbook=openpyxl.load_workbook(source,data_only=True)
    sheet=workbook["性能数据表"]
    group=formula=None; rows=[]
    for number in range(2,sheet.max_row+1):
        values=[sheet.cell(number,c).value for c in range(1,11)]
        if values[0] is not None:group=values[0]
        if values[1] is not None:formula=values[1]
        if group is None or formula is None or any(type(v) not in (int,float) or not math.isfinite(v) for v in values[2:]):
            raise ValueError("unexpected missing/nonfinite experimental row")
        for col in (1,2):
            if sheet.cell(number,col).value is None and not any(r.min_col<=col<=r.max_col and r.min_row<=number<=r.max_row for r in sheet.merged_cells.ranges):
                raise ValueError("blank group label outside a merged source cell")
        if not all(0<=v<=100 for v in values[3:]) or abs(math.fsum(values[4:])-100)>.011:
            raise ValueError("percentage bounds or selectivity closure failed")
        rows.append({"source_row":number,"group":group,"formula":formula,"temperature_c":values[2],
                     "conversion_pct":values[3],"selectivity_pct":values[5],"yield_pct":values[3]*values[5]/100,
                     "mode_II":int(group.startswith("B")),**parse_formula(formula)})
    workbook.close()
    if len(rows)!=114 or len({r["group"] for r in rows})!=21:
        raise ValueError("reviewed experimental design changed")
    return rows


def prepare(source,run):
    for name,digest in HASHES.items():
        if sha256_file(source/name)!=digest:raise ValueError("source differs from reviewed attachment: "+name)
    initialize_run(run,source/"CUMCM2021-B.pdf",[source/"附件1.xlsx",source/"附件2.xlsx"],"2021 B 统计建模复核","ethanol-2021b")
    records=extract(run/"sources/attachment_01__附件1.xlsx")
    write_json(run/"artifacts/observations.json",{"features":FEATURES,"rows":records,"source_hashes":HASHES})
    # The old solution has already been read; freeze choices before this reanalysis, not as a blind protocol.
    columns=list(range(len(FEATURES)))
    mean={"name":"training_mean","kind":"mean"}
    temp={"name":"temperature_quadratic","kind":"polynomial","degree":2,"columns":[0]}
    ridge=[{"name":f"ridge_d{degree}_a{alpha}","kind":"polynomial","degree":degree,"alpha":alpha,"columns":columns}
           for degree in (1,2) for alpha in (1.,10.,100.)]
    forest=[{"name":f"forest_leaf{leaf}","kind":"forest","trees":200,"leaf":leaf,"seed":20210904,"columns":columns}
            for leaf in (2,4)]
    families={"training_mean":[mean],"temperature_only":[temp],"ridge":ridge,"random_forest":forest,
              "nested_selection":[mean,temp,*ridge,*forest]}
    protocol={"features":FEATURES,"targets":["conversion_pct","selectivity_pct","yield_pct"],"families":families,
              "outer_folds":5,"inner_folds":4,"independent_unit":"catalyst combination, never individual temperature row",
              "physical_bounds":[0,100],"selection_metric":"root of mean per-group mean squared error",
              "permutation":"entire composition descriptor block exchanged among held-out groups, temperature retained; 20 seeds, predictive dependence only",
              "history_scope":"retrospective reanalysis after inspecting prior 2021 B solution; old files preserved; not a blind trial",
              "old_issues":["global target mean was labeled as a held-out mean baseline","independently permuting mass, total mass and ratio creates inconsistent descriptors","349.75 C is a grid-specific candidate, not a continuous maximum under strict T<350"],
              "observations_sha256":sha256_file(run/"artifacts/observations.json")}
    write_json(run/"artifacts/protocol.json",protocol)
    contract=read_json(run/"problem_contract.json")
    contract.update(problem_interpretation="Describe each formulation, compare grouped predictive models, distinguish observed candidates and open-domain optima, propose exactly five new experiments.",
        official_outputs=["21 formulation temperature relations and one-run stability", "formulation and temperature associations", "unrestricted and strict-below-350 candidate settings", "five experiments and reasons"],
        questions=[{"id":f"Q{i}","source_locator":f"PDF page 1 question ({i}), page 2 definitions and attachment notes", "objective":task,
                    "inputs":["frozen original attachments"],"outputs":[task],"decision_variables":["model/formulation/temperature/experiment design as applicable"],
                    "hard_constraints":["percentages 0-100; yield=X*S/100 in percentage units", "catalyst combination is the grouping unit", "low-temperature domain is T<350", "at most five additional experiments"],
                    "units":["Celsius", "percentage points", "mg", "ml/min"],"ambiguities":["no replicate variance or catalyst aging alignment in attachment 1", "instrument temperature resolution not specified"],
                    "termination_condition":"complete reported outputs and independent source/metric checks; no award or physical-global-optimum certification"}
                   for i,task in enumerate(["fit each formulation temperature relation and describe stability","estimate predictive associations with fair baselines","find observed and model-conditional candidates and open-boundary supremum","design five information-seeking experiments"],1)])
    write_json(run/"problem_contract.json",contract)
    audit=read_json(run/"source_audit.json")
    for row in audit["sources"]:
        if row["role"]=="problem":row["manual_review"]={"completed":True,"evidence":"Both original PDF pages read and rendered; four questions and percentage definitions checked. Agent reading, not human sign-off."}
    write_json(run/"source_audit.json",audit)
    print({"observations":len(records),"groups":21,"protocol_frozen":True},flush=True)


if __name__=="__main__":
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument("--source-dir",type=Path,required=True);parser.add_argument("--output",type=Path,required=True)
    args=parser.parse_args();prepare(args.source_dir,args.output)
